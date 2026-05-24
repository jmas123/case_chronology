"""Document text extraction.

One function per supported type. Each returns an ExtractedDocument with
raw_text (all pages joined with form-feed), page_count, and a warning
string when extraction is degraded (e.g. scanned PDF with no text layer).

raw_text uses '\\f' (form feed, ASCII 12) as the page separator. Page n
is the n-th segment when raw_text.split('\\f'). This is how Source
char_start/char_end offsets are interpreted in later phases.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

PAGE_SEP = "\f"


@dataclass(frozen=True)
class ExtractedDocument:
    raw_text: str
    page_count: int
    warning: str | None = None


def extract_pdf(path: Path) -> ExtractedDocument:
    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    raw_text = PAGE_SEP.join(pages)
    warning = None
    if not any(pages):
        warning = (
            "No text could be extracted from this PDF. It is likely a scanned "
            "image without a text layer. OCR is not supported in v1."
        )
    return ExtractedDocument(raw_text=raw_text, page_count=len(pages), warning=warning)


def extract_docx(path: Path) -> ExtractedDocument:
    doc = DocxDocument(str(path))
    paragraphs = [p.text for p in doc.paragraphs]
    raw_text = "\n".join(paragraphs).strip()
    # .docx has no reliable page break info from python-docx. Treat the whole
    # document as a single page. Documented in ROADMAP Phase 1.
    return ExtractedDocument(raw_text=raw_text, page_count=1)


def extract_txt(path: Path) -> ExtractedDocument:
    raw_text = path.read_text(encoding="utf-8", errors="replace").strip()
    return ExtractedDocument(raw_text=raw_text, page_count=1)


def extract(path: Path, ext: str) -> ExtractedDocument:
    """Dispatch on extension. Caller has already validated the extension."""
    if ext == "pdf":
        return extract_pdf(path)
    if ext == "docx":
        return extract_docx(path)
    if ext == "txt":
        return extract_txt(path)
    raise ValueError(f"Unsupported extension: {ext}")
